import pandas as pd
import numpy as np
import os
import re
from datetime import datetime

# ──────────────────────────────────────────────
# FIELD MAPPING
# ──────────────────────────────────────────────

# .dat column names (assigned manually — file has no header row)
DAT_COLUMNS = [
    "Scenario", "Year", "Period", "View",
    "Entity", "Value", "Account", "ICP",
    "Custom1", "Custom2", "Custom3", "Custom4",
    "Amount",
]

# Join keys: (HYP .dat column, BW Excel column)
JOIN_KEYS: list[tuple[str, str]] = [
    ("Scenario", "/BIC/SCNRBI"),
    ("Year",     "CALYEAR"),
    ("Period",   "MONTH"),
    ("Entity",   "/BIC/COMPBI"),
    ("Account",  "/BIC/ACCOUNTBI"),
    ("ICP",      "/BIC/ICPBI"),
    ("Custom1",  "/BIC/CUSTOM1BI"),
    ("Custom2",  "/BIC/CUSTOM2BI"),
    ("Custom3",  "/BIC/CUSTOM3BI"),
    ("Custom4",  "/BIC/CUSTOM4BI"),
]

# Numeric comparison: (HYP .dat column, BW Excel column)
NUMERIC_MAP: list[tuple[str, str]] = [
    ("Amount", "CUM_VALUE"),
]

# Columns not used in matching or comparison
HYP_SKIP  = ["View", "Value"]
BW_SKIP   = ["MTD_VALUE", "/BIC/CURBI", "/BIC/CURTYPE", "/BIC/ICPFLAG"]


# ──────────────────────────────────────────────
# FILE LOADING
# ──────────────────────────────────────────────

def pick_file(title: str) -> str:
    print(f"\n  {title}")
    print("  TIP: Drag the file from Finder into this terminal, then press ENTER.")
    path = input("  File path: ").strip().strip("'\"")
    if not path or not os.path.isfile(path):
        raise SystemExit(f"File not found: {path!r}")
    return path


def load_dat_file() -> pd.DataFrame:
    print("\n─────────────────────────────────────────")
    print("  STEP 1 — Load Hyperion .dat file")
    print("─────────────────────────────────────────")

    path = pick_file("Select the Hyperion .dat file")
    print(f"  File : {os.path.basename(path)}")

    # Read UTF-16, skip the '! DATA' first line, no header, semicolon delimited
    df = pd.read_csv(
        path,
        encoding="utf-16",
        sep=";",
        skiprows=1,       # skip '! DATA' line
        header=None,
        names=DAT_COLUMNS,
        dtype=str,
        on_bad_lines="warn",
    )
    df = df.apply(lambda c: c.str.strip() if c.dtype == object else c)
    df = df.dropna(how="all").reset_index(drop=True)

    print(f"  Rows loaded : {len(df):,}")
    print(f"  Sample (first 3 rows):")
    print(df.head(3).to_string(index=False))
    return df


def load_bw_file() -> pd.DataFrame:
    print("\n─────────────────────────────────────────")
    print("  STEP 2 — Load BW Excel backup file")
    print("─────────────────────────────────────────")

    path = pick_file("Select the BW Excel backup file")
    print(f"  File : {os.path.basename(path)}")

    xl = pd.ExcelFile(path)
    if len(xl.sheet_names) > 1:
        print(f"  Sheets: {xl.sheet_names}")
        sheet = input("  Sheet name (ENTER = first): ").strip() or xl.sheet_names[0]
    else:
        sheet = xl.sheet_names[0]

    df = pd.read_excel(path, sheet_name=sheet, dtype=str, header=0)
    df.columns = df.columns.str.strip()
    df = df.apply(lambda c: c.str.strip() if c.dtype == object else c)
    df = df.dropna(how="all").reset_index(drop=True)

    print(f"  Sheet : {sheet!r}  |  Rows : {len(df):,}")
    print(f"  Sample (first 3 rows):")
    bw_show = [b for _, b in JOIN_KEYS] + [b for _, b in NUMERIC_MAP]
    bw_show = [c for c in bw_show if c in df.columns]
    print(df[bw_show].head(3).to_string(index=False))
    return df


# ──────────────────────────────────────────────
# NORMALISATION
# ──────────────────────────────────────────────

MONTH_MAP = {
    "jan":"1","feb":"2","mar":"3","apr":"4","may":"5","jun":"6",
    "jul":"7","aug":"8","sep":"9","oct":"10","nov":"11","dec":"12",
}

def norm_period(s: pd.Series) -> pd.Series:
    def _f(v):
        v = str(v).strip()
        lo = v.lower()
        if lo in MONTH_MAP:
            return MONTH_MAP[lo]
        try:
            return str(int(float(v)))
        except Exception:
            return v
    return s.map(_f)


_NONE_RE = re.compile(r"^\[?(icp\s*)?none\]?$", re.IGNORECASE)

def norm_icp(s: pd.Series) -> pd.Series:
    def _f(v):
        v = str(v).strip()
        return "NONE" if _NONE_RE.match(v) else v.strip("[]").strip()
    return s.map(_f)


def normalise(df: pd.DataFrame, key_cols: list[str], source: str) -> pd.DataFrame:
    df = df.copy()
    for col in key_cols:
        if col not in df.columns:
            continue
        # Uppercase all join keys for case-insensitive matching
        df[col] = df[col].astype(str).str.strip().str.upper()
        if col in ("Period", "MONTH"):
            df[col] = norm_period(df[col])
        if col in ("ICP", "/BIC/ICPBI"):
            df[col] = norm_icp(df[col])
    print(f"  [{source}] Keys normalised (uppercase, period→number, ICP→NONE)")
    return df


# ──────────────────────────────────────────────
# COMPARISON
# ──────────────────────────────────────────────

def compare(hyp: pd.DataFrame, bw: pd.DataFrame) -> pd.DataFrame:
    hyp_key_cols = [h for h, _ in JOIN_KEYS]
    bw_key_cols  = [b for _, b in JOIN_KEYS]

    # Validate
    missing_hyp = [c for c in hyp_key_cols if c not in hyp.columns]
    missing_bw  = [c for c in bw_key_cols  if c not in bw.columns]
    if missing_hyp:
        print(f"  [ERROR] .dat file missing columns: {missing_hyp}")
        print(f"  .dat columns found: {list(hyp.columns)}")
        raise ValueError("Column mismatch in .dat file")
    if missing_bw:
        print(f"  [ERROR] BW Excel missing columns: {missing_bw}")
        print(f"  BW columns found: {list(bw.columns)}")
        raise ValueError("Column mismatch in BW file")

    # Normalise
    hyp_c = normalise(hyp, hyp_key_cols, "HYP")
    bw_c  = normalise(bw,  bw_key_cols,  "BW")

    # Rename BW columns to HYP names for merge
    bw_c = bw_c.rename(columns={b: h for h, b in JOIN_KEYS})
    bw_c = bw_c.rename(columns={b: f"{h}_BW" for h, b in NUMERIC_MAP})

    bw_keep = hyp_key_cols + [f"{h}_BW" for h, _ in NUMERIC_MAP]
    bw_c = bw_c[[c for c in bw_keep if c in bw_c.columns]]

    merged = hyp_c.merge(bw_c, on=hyp_key_cols, how="outer", indicator=True)

    results = []
    for hyp_col, bw_col_orig in NUMERIC_MAP:
        bw_col = f"{hyp_col}_BW"
        if hyp_col not in merged.columns or bw_col not in merged.columns:
            continue

        hyp_vals = pd.to_numeric(merged[hyp_col].astype(str).str.strip(), errors="coerce")
        bw_vals  = pd.to_numeric(merged[bw_col].astype(str).str.strip(),  errors="coerce")
        diff     = hyp_vals - bw_vals
        abs_diff = diff.abs()

        # Rows missing from one side, or where values differ (including sign-only)
        mismatch = (merged["_merge"] != "both") | (abs_diff > 1e-4)

        sub = merged.loc[mismatch, hyp_key_cols].copy()
        sub[f"HYP_{hyp_col}"]        = hyp_vals[mismatch].values
        sub[f"BW_{bw_col_orig}"]      = bw_vals[mismatch].values
        sub["Difference (HYP-BW)"]   = diff[mismatch].values
        sub["Abs Difference"]         = abs_diff[mismatch].values

        # Classify: sign-only difference vs real value difference
        merge_flag = merged.loc[mismatch, "_merge"].values
        hyp_m = hyp_vals[mismatch].values
        bw_m  = bw_vals[mismatch].values

        statuses = []
        for flag, h, b in zip(merge_flag, hyp_m, bw_m):
            if flag == "left_only":
                statuses.append("HYP only — not in BW")
            elif flag == "right_only":
                statuses.append("BW only — not in HYP")
            elif not np.isnan(h) and not np.isnan(b) and abs(abs(h) - abs(b)) <= 1e-4:
                statuses.append("Sign difference only")
            else:
                statuses.append("Value mismatch")
        sub["Status"] = statuses
        results.append(sub)

    return pd.concat(results, ignore_index=True) if results else pd.DataFrame()


# ──────────────────────────────────────────────
# OUTPUT
# ──────────────────────────────────────────────

def print_summary(result: pd.DataFrame) -> None:
    if result.empty:
        print("\n  All values match — no discrepancies found.")
        return

    hyp_only  = (result["Status"] == "HYP only — not in BW").sum()
    bw_only   = (result["Status"] == "BW only — not in HYP").sum()
    sign_diff = (result["Status"] == "Sign difference only").sum()
    val_diff  = (result["Status"] == "Value mismatch").sum()
    net       = result["Difference (HYP-BW)"].sum()

    print(f"\n  ┌──────────────────────────────────────┐")
    print(f"  │  COMPARISON SUMMARY                  │")
    print(f"  ├──────────────────────────────────────┤")
    print(f"  │  Total discrepancies  : {len(result):>9,}  │")
    print(f"  │  HYP only (no BW row) : {hyp_only:>9,}  │")
    print(f"  │  BW only (no HYP row) : {bw_only:>9,}  │")
    print(f"  │  Sign difference only : {sign_diff:>9,}  │")
    print(f"  │  Value mismatches     : {val_diff:>9,}  │")
    print(f"  │  Net diff (HYP-BW)   : {net:>9,.2f}  │")
    print(f"  └──────────────────────────────────────┘")


def export_results(result: pd.DataFrame) -> None:
    from datetime import datetime
    from openpyxl.styles import PatternFill, Font

    print_summary(result)
    if result.empty:
        return

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path  = f"/Users/ksen/Downloads/HYP_BW_differences_{timestamp}.xlsx"

    STATUS_COLORS = {
        "HYP only — not in BW":  "FFA500",   # orange
        "BW only — not in HYP":  "ADD8E6",   # light blue
        "Sign difference only":   "FFFF99",   # yellow
        "Value mismatch":         "FF9999",   # red
    }

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        result.to_excel(writer, sheet_name="Differences", index=False)
        ws = writer.sheets["Differences"]

        # Auto-fit columns
        for col_cells in ws.columns:
            width = max(len(str(c.value or "")) for c in col_cells) + 2
            ws.column_dimensions[col_cells[0].column_letter].width = min(width, 45)

        # Bold header
        for cell in ws[1]:
            cell.font = Font(bold=True)

        # Colour rows by status
        status_col_idx = result.columns.get_loc("Status") + 1
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            status_val = row[status_col_idx - 1].value
            colour = STATUS_COLORS.get(status_val)
            if colour:
                fill = PatternFill(fill_type="solid", fgColor=colour)
                for cell in row:
                    cell.fill = fill

    print(f"\n  Saved → {out_path}")


# ──────────────────────────────────────────────
# MAIN
# ──────────────────────────────────────────────

def main():
    print("=" * 50)
    print("  HYP_check — Hyperion vs BW Comparison")
    print("=" * 50)

    hyp_df = load_dat_file()
    bw_df  = load_bw_file()

    print("\n─────────────────────────────────────────")
    print("  STEP 3 — Comparing")
    print("─────────────────────────────────────────")

    try:
        result = compare(hyp_df, bw_df)
        export_results(result)
    except ValueError as e:
        print(f"  [STOPPED] {e}")

    print("\nDone.")


if __name__ == "__main__":
    main()
